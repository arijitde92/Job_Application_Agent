"""Tests for tailoring from a pasted/uploaded job description (issue #10)."""

import json
import os

import pytest
from fastapi.testclient import TestClient
from pydantic import ValidationError

from app.api.deps import get_current_user, get_db
from app.main import app
from app.schemas import TailorJobRequest
from app.services import crew_runner
from app.services.extractors import text_job_extractor
from app.services.extractors.text_job_extractor import (
    JobDescriptionParsingError,
    extract_job_description_file_text,
    extract_job_details_from_text,
)

JD_TEXT = (
    "Senior Backend Engineer — Acme Robotics (Berlin, Hybrid)\n\n"
    "We are looking for a backend engineer to build our fleet telemetry "
    "platform in Python and FastAPI.\n\nRequirements:\n- 5+ years of Python\n"
    "- Experience with PostgreSQL and Kubernetes\n"
)

# A real job-board posting, pasted as-is: icon labels ("clock icon"), a header
# location (Kolkata) that disagrees with the body (Hyderabad), and em-dashes
# lost to double spaces. The job description textbox must cope with this.
HYLAND_JD_TEXT = """\
Key Highlights
clock icon0 to 4 Yrs
subtract iconNot Disclosed
location icon
Kolkata
Job Description
Overview
Senior Software Engineer (AI Enablement)
   Location: Hyderabad    Work Arrangement: Hybrid  2 days per week in the office




About the Role
We're looking for a technically strong and driven Senior Software Engineer to take ownership of complex engineering challenges and deliver high-quality software that powers Hyland's enterprise platforms. In this role, you'll apply solid software engineering principles to the design, development, maintenance, testing, and evaluation of software  ensuring timely delivery within release timelines and quality guidelines.


You'll contribute meaningfully to architectural discussions, mentor fellow engineers, and help shape team standards and engineering practices. If you're passionate about writing clean, impactful code, solving hard problems, and staying ahead of the curve on emerging technologies  including AI-assisted and agentic AI development  we'd love to hear from you.




Your Role Responsibilities  Here's What You'll Do.
Design and develop complex, high-quality code based on functional specifications  translating and analyzing software requirements into design specifications and ensuring adherence to established standards, algorithms, and logic design.
Perform complex peer code reviews and analysis, providing insight on the broader impact of software changes while influencing and guiding other team members on technical direction and best practices.
Create and apply automated tests and test principles  including unit tests  to software changes; contribute to the implementation of a delivery pipeline inclusive of test automation, security, and performance validation.
Incorporate business value into engineering decisions, influence team standards and processes, and lead planning sessions, work estimation, solution demos, and implementation design discussions.
Research and resolve complex escalations for production issues; implement development standards to ensure compliance with product and industry regulations and recommend internal process improvements and product documentation enhancements.
Mentor, coach, and provide feedback to team members; leverage AI-assisted and agentic AI development tools to drive productivity, code quality, and innovation across the engineering team.


Technology Tools
Java
C / C++
Python
Open Source Tools & Platforms
CI/CD Pipelines & Build Environments
AI-Assisted Development Tools (GitHub Copilot, Cursor, Claude)
Agentic AI Frameworks & LLM APIs


Role Essentials
Bachelor's degree or equivalent experience; significant experience with one or more general-purpose programming languages  including Java, C/C++, C#, Python, or JavaScript  with demonstrated strength in data structures, algorithms, and software design.
Experience with continuous software delivery, build environments, delivery pipelines, test automation, and continuous integration tools in Windows/Linux development environments using open-source tools and platforms.
Strong understanding of software application testing tools, methodologies, and process frameworks; ability to create and apply automated tests  including unit tests  to software changes effectively.
Strong critical thinking and problem-solving skills; self-motivated with the ability to manage projects to completion with minimal oversight; demonstrated ability to influence, motivate, and mobilize team members and business partners.
Strong oral and written communication skills with the ability to interact with others with discretion and tact; effective at building rapport, gaining consensus, and translating goals into new ideas and design solutions across all levels of the organization.


What We'd Like to See
Hands-on experience with AI-assisted development tools (GitHub Copilot, Cursor, Claude) and agentic AI frameworks for autonomous workflow orchestration and intelligent code generation.
Experience with agentic AI patterns  including multi-agent orchestration, tool-use, memory management, and autonomous task execution  within enterprise software systems.
Exposure to LLM APIs, prompt engineering, RAG architectures, or AI-powered automation capabilities integrated into software delivery workflows.
Experience working with cloud platforms (AWS or Azure) and containerized environments (Docker, Kubernetes) for scalable, cloud-native application development.
A passion for staying current with emerging technologies and a track record of continuous learning, applying new technical concepts to improve product quality and team capability.


About Hyland
Hyland is the pioneer of the Content Innovation Cloud, delivering ubiquitous enterprise intelligence to organizations with solutions that unlock actionable insights and drive automation.


Trusted by thousands of organizations worldwide, including many of the Fortune 100, Hyland's solutions create the foundation for a connected, agentic enterprise, where teams harness the power of AI to redefine how they operate and engage with those they serve. For .
Other Details
Industry
IT Services & Consulting
Recruiter Details
Hyland Software
Job Tags
java, python
Job Type
Full time
"""


# ── TailorJobRequest: exactly one job source ─────────────────────────────────

def test_request_accepts_linkedin_url_only():
    req = TailorJobRequest(linkedin_job_url="https://www.linkedin.com/jobs/view/1/", resume_id=1)
    assert req.linkedin_job_url == "https://www.linkedin.com/jobs/view/1/"
    assert req.job_description_text is None


def test_request_accepts_job_description_only_and_strips_it():
    req = TailorJobRequest(job_description_text=f"  {JD_TEXT}  ", resume_id=1)
    assert req.job_description_text == JD_TEXT.strip()
    assert req.linkedin_job_url is None


def test_request_treats_blank_strings_as_missing():
    req = TailorJobRequest(linkedin_job_url="   ", job_description_text=JD_TEXT, resume_id=1)
    assert req.linkedin_job_url is None


@pytest.mark.parametrize("kwargs, message", [
    ({}, "is required"),
    ({"linkedin_job_url": "", "job_description_text": "  "}, "is required"),
    ({"linkedin_job_url": "https://www.linkedin.com/jobs/view/1/",
      "job_description_text": JD_TEXT}, "not both"),
    ({"job_description_text": "Backend engineer"}, "too short"),
])
def test_request_rejects_invalid_job_sources(kwargs, message):
    with pytest.raises(ValidationError, match=message):
        TailorJobRequest(resume_id=1, **kwargs)


def test_request_accepts_real_job_board_posting():
    req = TailorJobRequest(job_description_text=HYLAND_JD_TEXT, resume_id=1)
    assert req.job_description_text == HYLAND_JD_TEXT.strip()


def test_request_rejects_overlong_job_description():
    with pytest.raises(ValidationError):
        TailorJobRequest(
            job_description_text="x" * (text_job_extractor.MAX_JOB_DESCRIPTION_CHARS + 1),
            resume_id=1,
        )


# ── Job description file → text ──────────────────────────────────────────────

@pytest.mark.parametrize("suffix", [".txt", ".md"])
def test_extract_plain_text_files(tmp_path, suffix):
    path = tmp_path / f"jd{suffix}"
    path.write_text(JD_TEXT, encoding="utf-8")
    assert extract_job_description_file_text(str(path)) == JD_TEXT.strip()


def test_extract_docx_file(tmp_path):
    from docx import Document

    document = Document()
    for line in JD_TEXT.splitlines():
        document.add_paragraph(line)
    path = tmp_path / "jd.docx"
    document.save(path)

    text = extract_job_description_file_text(str(path))
    assert "Senior Backend Engineer" in text
    assert "5+ years of Python" in text


@pytest.mark.parametrize("suffix", [".txt", ".docx"])
def test_extract_real_job_board_posting_file(tmp_path, suffix):
    path = tmp_path / f"hyland{suffix}"
    if suffix == ".docx":
        from docx import Document

        document = Document()
        for line in HYLAND_JD_TEXT.splitlines():
            document.add_paragraph(line)
        document.save(path)
    else:
        path.write_text(HYLAND_JD_TEXT, encoding="utf-8")

    text = extract_job_description_file_text(str(path))

    if suffix == ".txt":
        assert text == HYLAND_JD_TEXT.strip()
    else:  # python-docx extraction drops the blank paragraphs
        non_blank = [line for line in HYLAND_JD_TEXT.splitlines() if line.strip()]
        assert text.splitlines() == non_blank


def test_extract_rejects_unsupported_type(tmp_path):
    path = tmp_path / "jd.html"
    path.write_text(JD_TEXT)
    with pytest.raises(JobDescriptionParsingError, match="Unsupported"):
        extract_job_description_file_text(str(path))


def test_extract_rejects_near_empty_file(tmp_path):
    path = tmp_path / "jd.txt"
    path.write_text("   Engineer  \n")
    with pytest.raises(JobDescriptionParsingError, match="no meaningful text"):
        extract_job_description_file_text(str(path))


def test_extract_wraps_parser_errors(tmp_path):
    path = tmp_path / "jd.docx"
    path.write_bytes(b"not really a docx")
    with pytest.raises(JobDescriptionParsingError):
        extract_job_description_file_text(str(path))


# ── Job description text → JobDetails ────────────────────────────────────────

def _fake_llm(reply, jd_text=JD_TEXT):
    def call(prompt):
        assert jd_text.strip() in prompt
        if isinstance(reply, Exception):
            raise reply
        return reply
    return call


def test_details_from_text_uses_llm_metadata_and_keeps_text_verbatim(monkeypatch):
    reply = "```json\n" + json.dumps({
        "job_name": "Senior Backend Engineer",
        "company_name": "Acme Robotics",
        "location": "Berlin",
        "workplace_type": "Hybrid",
        "seniority_level": "",
        "employment_type": None,
        "requirements": ["5+ years of Python", "PostgreSQL and Kubernetes"],
    }) + "\n```"
    monkeypatch.setattr(text_job_extractor, "_call_metadata_llm", _fake_llm(reply))

    details = extract_job_details_from_text(JD_TEXT)

    assert details.job_description == JD_TEXT.strip()
    assert details.url == ""
    assert details.job_name == "Senior Backend Engineer"
    assert details.company_name == "Acme Robotics"
    assert details.workplace_type == "Hybrid"
    # Blank / null / missing values are normalised to "N/A".
    assert details.seniority_level == "N/A"
    assert details.employment_type == "N/A"
    assert details.industry == "N/A"
    assert details.requirements == "5+ years of Python\nPostgreSQL and Kubernetes"


@pytest.mark.parametrize("llm_value, expected", [
    ("Mid-Senior level", "Mid-Senior level"),
    ("mid senior  LEVEL", "Mid-Senior level"),    # case / hyphen / spacing
    ("entry-level", "Entry level"),
    ("Internship", "Internship"),
    ("Executive", "Executive"),
    ("Senior Software Engineer", "N/A"),          # the job title echoed back
    ("Senior", "N/A"),                            # not a standard level
    ("N/A", "N/A"),
])
def test_details_from_text_normalises_seniority_level(monkeypatch, llm_value, expected):
    reply = json.dumps({"seniority_level": llm_value})
    monkeypatch.setattr(text_job_extractor, "_call_metadata_llm", _fake_llm(reply))

    assert extract_job_details_from_text(JD_TEXT).seniority_level == expected


@pytest.mark.parametrize("reply", [
    "Sorry, I cannot help with that.",   # not JSON
    "[1, 2, 3]",                         # JSON, but not an object
    RuntimeError("Gemini unavailable"),  # the LLM call itself fails
])
def test_details_from_text_falls_back_to_na_metadata(monkeypatch, reply):
    monkeypatch.setattr(text_job_extractor, "_call_metadata_llm", _fake_llm(reply))

    details = extract_job_details_from_text(JD_TEXT)

    assert details.job_description == JD_TEXT.strip()
    assert details.job_name == "N/A"
    assert details.company_name == "N/A"


def test_details_from_real_job_board_posting_keeps_text_verbatim(monkeypatch):
    reply = json.dumps({
        "job_name": "Senior Software Engineer (AI Enablement)",
        "company_name": "Hyland",
        "location": "Hyderabad",
    })
    monkeypatch.setattr(
        text_job_extractor, "_call_metadata_llm", _fake_llm(reply, HYLAND_JD_TEXT)
    )

    details = extract_job_details_from_text(HYLAND_JD_TEXT)

    # The whole posting — noise lines and double spaces included — reaches the
    # crew unchanged; only the metadata comes from the LLM.
    assert details.job_description == HYLAND_JD_TEXT.strip()
    assert details.job_name == "Senior Software Engineer (AI Enablement)"
    assert details.company_name == "Hyland"


@pytest.mark.skipif(
    os.environ.get("RUN_LIVE_LLM_TESTS") != "1",
    reason="live Gemini call — set RUN_LIVE_LLM_TESTS=1 (needs GEMINI_API_KEY)",
)
def test_live_llm_extracts_real_job_board_posting():
    from dotenv import load_dotenv

    load_dotenv()
    details = extract_job_details_from_text(HYLAND_JD_TEXT)

    assert details.job_description == HYLAND_JD_TEXT.strip()
    assert "Senior Software Engineer" in details.job_name
    assert "Hyland" in details.company_name
    # The body's "Location: Hyderabad" wins over the board's "Kolkata" header.
    assert "Hyderabad" in details.location
    assert details.workplace_type == "Hybrid"
    # Derived from "Senior" in the title — it outranks the board's "0 to 4 Yrs".
    assert details.seniority_level == "Mid-Senior level"
    assert details.employment_type.lower().replace("-", " ") == "full time"
    assert details.industry == "IT Services & Consulting"
    assert "Bachelor's degree" in details.requirements


# ── API ──────────────────────────────────────────────────────────────────────

class _User:
    id = 1
    email = "applicant@example.com"
    first_name = "Test"
    last_name = "User"


@pytest.fixture
def client():
    async def no_db():
        yield None

    app.dependency_overrides[get_current_user] = lambda: _User()
    app.dependency_overrides[get_db] = no_db
    yield TestClient(app)
    app.dependency_overrides.clear()


def test_extract_endpoint_returns_file_text(client):
    resp = client.post(
        "/api/jobs/description/extract",
        files={"file": ("posting.txt", JD_TEXT.encode(), "text/plain")},
    )
    assert resp.status_code == 200
    assert resp.json() == {"filename": "posting.txt", "text": JD_TEXT.strip()}


def test_extract_endpoint_returns_real_job_board_posting(client):
    resp = client.post(
        "/api/jobs/description/extract",
        files={"file": ("hyland.txt", HYLAND_JD_TEXT.encode(), "text/plain")},
    )
    assert resp.status_code == 200
    assert resp.json()["text"] == HYLAND_JD_TEXT.strip()


def test_extract_endpoint_rejects_unsupported_type(client):
    resp = client.post(
        "/api/jobs/description/extract",
        files={"file": ("posting.html", JD_TEXT.encode(), "text/html")},
    )
    assert resp.status_code == 400


def test_extract_endpoint_rejects_empty_file(client):
    resp = client.post(
        "/api/jobs/description/extract",
        files={"file": ("posting.txt", b"   ", "text/plain")},
    )
    assert resp.status_code == 422
    assert "posting.txt" in resp.json()["detail"]


def test_extract_endpoint_requires_auth():
    resp = TestClient(app).post(
        "/api/jobs/description/extract",
        files={"file": ("posting.txt", JD_TEXT.encode(), "text/plain")},
    )
    assert resp.status_code in (401, 403)


def test_tailor_endpoint_rejects_both_job_sources(client):
    resp = client.post("/api/jobs/tailor", json={
        "linkedin_job_url": "https://www.linkedin.com/jobs/view/1/",
        "job_description_text": JD_TEXT,
        "resume_id": 1,
    })
    assert resp.status_code == 422
    assert "not both" in json.dumps(resp.json())


# ── crew_runner: job source routing ──────────────────────────────────────────

def _stop_after_job_extraction(monkeypatch, calls):
    """Record which extractor runs; an unsupported resume type then aborts the
    run at the resume-parsing step, before any crew is built."""
    from app.services.extractors import linkedin_extractor
    from app.services.extractors.linkedin_extractor import JobDetails

    monkeypatch.setattr(
        text_job_extractor, "extract_job_details_from_text",
        lambda text: calls.append(("text", text)) or JobDetails(url="", job_description=text),
    )
    monkeypatch.setattr(
        linkedin_extractor, "extract_linkedin_job_details",
        lambda url: calls.append(("url", url)) or JobDetails(url=url),
    )


def test_crew_runner_uses_text_extractor_for_job_description(monkeypatch):
    calls = []
    _stop_after_job_extraction(monkeypatch, calls)
    with pytest.raises(RuntimeError, match="Resume parsing failed"):
        crew_runner._run_crew_sync(
            job_id=1, user_id=1, user_email="a@example.com",
            resume_bytes=b"x", resume_filename="resume.unsupported",
            job_description_text=JD_TEXT,
        )
    assert calls == [("text", JD_TEXT)]


def test_crew_runner_uses_linkedin_extractor_for_url(monkeypatch):
    calls = []
    _stop_after_job_extraction(monkeypatch, calls)
    url = "https://www.linkedin.com/jobs/view/1/"
    with pytest.raises(RuntimeError, match="Resume parsing failed"):
        crew_runner._run_crew_sync(
            job_id=1, user_id=1, user_email="a@example.com",
            resume_bytes=b"x", resume_filename="resume.unsupported", job_url=url,
        )
    assert calls == [("url", url)]
