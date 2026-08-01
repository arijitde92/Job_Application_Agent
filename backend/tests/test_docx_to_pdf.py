"""
Tests for the LibreOffice-based docx → PDF converter.

The conversion itself only runs where LibreOffice is installed (it is a
documented system dependency, not a Python one) — the round-trip test skips
cleanly elsewhere. The missing-binary behavior is tested everywhere, because
returning ``None`` instead of raising is what lets a crew run ship the .docx
without a PDF preview when LibreOffice is absent.
"""

import shutil

import pytest

import app.services.docx_to_pdf as docx_to_pdf
from app.services.docx_to_pdf import convert_docx_to_pdf

_HAS_LIBREOFFICE = bool(shutil.which("soffice") or shutil.which("libreoffice"))


def test_missing_libreoffice_returns_none(tmp_path, monkeypatch):
    monkeypatch.setattr(docx_to_pdf.shutil, "which", lambda _name: None)
    docx = tmp_path / "resume.docx"
    docx.write_bytes(b"PK\x03\x04irrelevant")
    assert convert_docx_to_pdf(str(docx)) is None


@pytest.mark.skipif(not _HAS_LIBREOFFICE, reason="LibreOffice not installed")
def test_roundtrip_produces_pdf(tmp_path):
    from docx import Document

    docx_path = tmp_path / "resume.docx"
    doc = Document()
    doc.add_heading("Ada Lovelace", level=1)
    doc.add_paragraph("Pioneering engineer with analytical experience.")
    doc.save(str(docx_path))

    pdf_path = convert_docx_to_pdf(str(docx_path), outdir=str(tmp_path))

    assert pdf_path is not None
    with open(pdf_path, "rb") as f:
        assert f.read(5) == b"%PDF-"


@pytest.mark.skipif(not _HAS_LIBREOFFICE, reason="LibreOffice not installed")
def test_corrupt_docx_returns_none(tmp_path):
    docx_path = tmp_path / "broken.docx"
    docx_path.write_bytes(b"this is not a zip archive at all")
    # LibreOffice either refuses (no output) or errors — both must yield None
    # or a real PDF; it must never raise.
    result = convert_docx_to_pdf(str(docx_path), outdir=str(tmp_path))
    assert result is None or result.endswith(".pdf")
